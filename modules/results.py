from datetime import datetime
import json
import csv
import io
from typing import Dict, List, Any, Optional

from modules.database import (
    db,
    Vote,
    Candidate,
    Position,
    Election,
    Student,
    Archive,
    AuditLog,
)


def _log_action(user: str, action: str):
    log = AuditLog(user=user, action=action, created_at=datetime.utcnow())
    db.session.add(log)
    db.session.commit()


def count_votes(election_id: Optional[int] = None) -> int:
    q = Vote.query
    if election_id:
        q = q.filter(Vote.election_id == election_id)
    return q.count()


def get_results(election_id: int) -> Dict[str, Any]:
    positions = Position.query.filter(Position.election_id == election_id).all()
    results = {}
    for pos in positions:
        results[pos.id] = get_results_by_position(election_id, pos.id)
    return results


def get_results_by_position(election_id: int, position_id: int) -> Dict[str, Any]:
    candidates = Candidate.query.filter(
        Candidate.election_id == election_id,
        Candidate.position_id == position_id,
        Candidate.status != "Deleted",
    ).all()

    total_valid_votes = (
        db.session.query(db.func.count(Vote.id))
        .filter(Vote.election_id == election_id)
        .join(Candidate, Candidate.id == Vote.candidate_id)
        .filter(Candidate.position_id == position_id)
        .scalar()
    )

    rows = []
    for c in candidates:
        votes = (
            Vote.query.filter(
                Vote.election_id == election_id, Vote.candidate_id == c.id
            ).count()
        )
        pct = calculate_percentages(votes, total_valid_votes)
        rows.append({
            "candidate_id": c.id,
            "name": c.name or c.full_name,
            "votes": votes,
            "percentage": pct,
            "photo": c.photo,
            "manifesto": c.manifesto,
        })

    # ranking
    rows.sort(key=lambda r: r["votes"], reverse=True)
    for idx, r in enumerate(rows, start=1):
        r["rank"] = idx

    position = Position.query.get(position_id)
    maximum_winners = position.maximum_winners if position else 1

    winners = [r for r in rows[:maximum_winners]]

    return {
        "position_id": position_id,
        "position_name": position.name if position else None,
        "total_valid_votes": int(total_valid_votes or 0),
        "candidates": rows,
        "winners": winners,
    }


def calculate_percentages(candidate_votes: int, total_valid_votes: int) -> float:
    try:
        if not total_valid_votes:
            return 0.0
        return round((candidate_votes / total_valid_votes) * 100, 2)
    except Exception:
        return 0.0


def calculate_participation(election_id: int) -> Dict[str, Any]:
    registered = Student.query.count()
    voted_students = (
        db.session.query(db.func.count(db.distinct(Vote.student_id)))
        .filter(Vote.election_id == election_id)
        .scalar()
    )
    remaining = max(0, registered - int(voted_students or 0))
    participation = (
        round((int(voted_students or 0) / registered) * 100, 2) if registered else 0.0
    )
    return {
        "registered": registered,
        "voted": int(voted_students or 0),
        "remaining": remaining,
        "participation_percent": participation,
    }


def determine_winners(election_id: int) -> Dict[int, List[Dict[str, Any]]]:
    positions = Position.query.filter(Position.election_id == election_id).all()
    winners = {}
    for pos in positions:
        res = get_results_by_position(election_id, pos.id)
        winners[pos.id] = res.get("winners", [])
    return winners


def detect_ties(election_id: int) -> Dict[int, Any]:
    ties = {}
    positions = Position.query.filter(Position.election_id == election_id).all()
    for pos in positions:
        res = get_results_by_position(election_id, pos.id)
        candidates = res.get("candidates", [])
        if not candidates:
            continue
        top_votes = candidates[0]["votes"]
        tied = [c for c in candidates if c["votes"] == top_votes]
        if len(tied) > 1:
            ties[pos.id] = {
                "position_name": pos.name,
                "tied_candidates": tied,
            }
    return ties


def generate_dashboard_statistics(election_id: int) -> Dict[str, Any]:
    election = Election.query.get(election_id)
    total_votes = count_votes(election_id)
    participation = calculate_participation(election_id)
    winners = determine_winners(election_id)
    positions_completed = sum(1 for p in Position.query.filter(Position.election_id == election_id).all() if get_results_by_position(election_id, p.id)["total_valid_votes"] > 0)

    return {
        "election_id": election_id,
        "election_title": election.title if election else None,
        "total_votes": total_votes,
        "participation": participation,
        "winners_count": sum(len(v) for v in winners.values()),
        "positions_completed": positions_completed,
        "last_updated": datetime.utcnow().isoformat(),
    }


def publish_results(election_id: int, user: str) -> bool:
    election = Election.query.get(election_id)
    if not election:
        return False
    election.status = "Published"
    election.show_live_results = True
    db.session.commit()
    _log_action(user, f"Published results for election {election_id}")
    return True


def unpublish_results(election_id: int, user: str) -> bool:
    election = Election.query.get(election_id)
    if not election:
        return False
    election.status = "Closed"
    election.show_live_results = False
    db.session.commit()
    _log_action(user, f"Unpublished results for election {election_id}")
    return True


def archive_results(election_id: int, user: str) -> bool:
    results = get_results(election_id)
    election = Election.query.get(election_id)
    archive = Archive(election_name=(election.title if election else f"election-{election_id}"), results=json.dumps(results), created_at=datetime.utcnow())
    db.session.add(archive)
    db.session.commit()
    _log_action(user, f"Archived results for election {election_id}")
    return True


def export_csv(election_id: int) -> str:
    """Return CSV content as string for the given election."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Position", "Candidate", "Votes", "Percentage"]) 
    positions = Position.query.filter(Position.election_id == election_id).all()
    for pos in positions:
        res = get_results_by_position(election_id, pos.id)
        for c in res.get("candidates", []):
            writer.writerow([pos.name, c.get("name"), c.get("votes"), c.get("percentage")])
    return output.getvalue()


def export_excel(election_id: int, file_path: str) -> bool:
    try:
        import pandas as pd

        rows = []
        positions = Position.query.filter(Position.election_id == election_id).all()
        for pos in positions:
            res = get_results_by_position(election_id, pos.id)
            for c in res.get("candidates", []):
                rows.append({
                    "position": pos.name,
                    "candidate": c.get("name"),
                    "votes": c.get("votes"),
                    "percentage": c.get("percentage"),
                })
        df = pd.DataFrame(rows)
        df.to_excel(file_path, index=False)
        _log_action("system", f"Exported election {election_id} to excel {file_path}")
        return True
    except Exception:
        return False


def export_pdf(election_id: int, file_path: str) -> bool:
    # Placeholder: PDF export requires reportlab or other library.
    return False


def get_candidate_statistics(candidate_id: int) -> Dict[str, Any]:
    c = Candidate.query.get(candidate_id)
    if not c:
        return {}
    votes = Vote.query.filter(Vote.candidate_id == candidate_id).count()
    return {
        "candidate_id": c.id,
        "name": c.name or c.full_name,
        "votes": votes,
        "manifesto": c.manifesto,
    }


def get_position_statistics(election_id: int, position_id: int) -> Dict[str, Any]:
    return get_results_by_position(election_id, position_id)


def get_live_statistics(election_id: int) -> Dict[str, Any]:
    # lightweight summary for live dashboards
    return generate_dashboard_statistics(election_id)


def verify_result_integrity(election_id: int) -> Dict[str, Any]:
    issues = []
    total_votes = count_votes(election_id)
    per_position_sum = 0
    positions = Position.query.filter(Position.election_id == election_id).all()
    for pos in positions:
        res = get_results_by_position(election_id, pos.id)
        per_position_sum += res.get("total_valid_votes", 0)

        # duplicate votes (student voted multiple times for same position)
        dup_q = (
            db.session.query(Vote.student_id, db.func.count(Vote.id).label("cnt"))
            .join(Candidate, Candidate.id == Vote.candidate_id)
            .filter(Candidate.position_id == pos.id, Vote.election_id == election_id)
            .group_by(Vote.student_id)
            .having(db.func.count(Vote.id) > 1)
        )
        duplicates = dup_q.all()
        if duplicates:
            issues.append({"position": pos.id, "duplicates": len(duplicates)})

    if per_position_sum != total_votes:
        issues.append({"mismatch": True, "total_votes": total_votes, "per_position_sum": per_position_sum})

    return {"ok": len(issues) == 0, "issues": issues}
from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    Response,
    session,
    jsonify
)

from modules.database import (
    db,
    Vote,
    Candidate,
    Position,
    Election,
    Student,
    AuditLog,
    Archive,
    SystemSetting
)

from modules.security import (
    admin_required,
    is_admin
)

import csv
import io
import json
from datetime import datetime
from collections import defaultdict


results = Blueprint(
    "results",
    __name__
)


def count_votes(election_id=None, position_id=None):
    query = Vote.query
    
    if election_id:
        query = query.filter(Vote.election_id == election_id)
    if position_id:
        query = query.join(Candidate).join(Position).filter(Position.id == position_id)
    
    total_votes = query.count()
    
    vote_counts = defaultdict(int)
    candidates_in_query = query.with_entities(Vote.candidate_id).distinct().all()
    
    for (candidate_id,) in candidates_in_query:
        vote_counts[candidate_id] = Vote.query.filter(
            Vote.candidate_id == candidate_id
        ).count()
    
    return dict(vote_counts) if not election_id else {
        k: v for k, v in vote_counts.items() 
        if k in [c.id for c in Candidate.query.filter_by(election_id=election_id).all()]
    }


def get_results(election_id=None):
    return get_results_by_position(election_id=election_id)


def get_results_by_position(election_id=None, position_id=None):
    election_filter = Election.query
    if election_id:
        election_filter = election_filter.filter(Election.id == election_id)

    elections = election_filter.order_by(Election.id.desc()).all()
    payload = []

    for election in elections:
        positions = Position.query.filter_by(election_id=election.id)
        if position_id:
            positions = positions.filter(Position.id == position_id)
        positions = positions.order_by(Position.display_order.asc(), Position.id.asc()).all()

        for position in positions:
            candidates = Candidate.query.filter_by(position_id=position.id).order_by(Candidate.display_order.asc(), Candidate.id.asc()).all()
            candidate_results = []
            total_valid_votes = 0

            for candidate in candidates:
                votes = Vote.query.filter_by(candidate_id=candidate.id, election_id=election.id).count()
                total_valid_votes += votes
                candidate_results.append({
                    "candidate_id": candidate.id,
                    "candidate_name": candidate.full_name or candidate.name,
                    "candidate_photo": candidate.photo,
                    "votes": votes,
                    "percentage": 0.0,
                    "manifesto": candidate.manifesto if hasattr(candidate, 'manifesto') else None,
                    "biography": candidate.biography if hasattr(candidate, 'biography') else None,
                    "slogan": candidate.slogan if hasattr(candidate, 'slogan') else None,
                })

            for item in candidate_results:
                item["percentage"] = calculate_percentages(item["votes"], total_valid_votes)

            candidate_results.sort(key=lambda item: (-item["votes"], item["candidate_name"]))
            for rank, item in enumerate(candidate_results, start=1):
                item["rank"] = rank
                item["winner_badge"] = rank == 1

            winner = candidate_results[0] if candidate_results else None
            payload.append({
                "election_id": election.id,
                "election_title": election.title,
                "election_status": election.status,
                "election_year": election.election_year,
                "position_id": position.id,
                "position_name": position.name,
                "position_code": position.position_code if hasattr(position, 'position_code') else None,
                "description": position.description if hasattr(position, 'description') else None,
                "candidates": candidate_results,
                "winner": winner,
                "total_votes": total_valid_votes,
            })

    return payload


def calculate_percentages(votes, total_valid_votes):
    if not total_valid_votes:
        return 0.0
    return round((votes / total_valid_votes) * 100, 1)


def calculate_participation(election_id=None):
    election_filter = Election.query
    if election_id:
        election_filter = election_filter.filter(Election.id == election_id)
    elections = election_filter.all()
    participation = []

    for election in elections:
        base_time = election.created_at or datetime.utcnow()
        eligible_query = Student.query.filter(
            Student.status == "Active",
            Student.created_at >= base_time,
        )

        if not eligible_query.first():
            eligible_query = Student.query.filter(Student.status == "Active")

        eligible_students = eligible_query.count()
        voted_students = (
            db.session.query(Vote.student_id)
            .filter(Vote.election_id == election.id)
            .distinct()
            .count()
        )

        if voted_students > eligible_students:
            eligible_students = voted_students

        not_voted_students = max(eligible_students - voted_students, 0)
        participation_rate = round((voted_students / eligible_students * 100), 1) if eligible_students else 0.0

        participation.append({
            "election_id": election.id,
            "election_title": election.title,
            "registered_students": eligible_students,
            "students_voted": voted_students,
            "students_remaining": not_voted_students,
            "participation_rate": participation_rate,
        })

    return participation


def determine_winners(election_id=None, position_id=None):
    results_data = get_results_by_position(election_id=election_id, position_id=position_id)
    winners = []
    for item in results_data:
        winner = item["winner"]
        candidates = item.get("candidates", [])
        runner_up = candidates[1] if len(candidates) > 1 else None
        third = candidates[2] if len(candidates) > 2 else None
        
        winners.append({
            "election_id": item["election_id"],
            "election_title": item["election_title"],
            "election_year": item.get("election_year"),
            "position_id": item["position_id"],
            "position_name": item["position_name"],
            "winner": winner,
            "candidate_name": winner["candidate_name"] if winner else "No winner",
            "votes": winner["votes"] if winner else 0,
            "percentage": winner["percentage"] if winner else 0,
            "runner_up": runner_up,
            "third": third,
        })
    return winners


def detect_ties(election_id=None, position_id=None):
    results_data = get_results_by_position(election_id=election_id, position_id=position_id)
    ties = []
    
    for item in results_data:
        candidates = item.get("candidates", [])
        if not candidates:
            continue
        
        max_votes = candidates[0]["votes"]
        tied_candidates = [c for c in candidates if c["votes"] == max_votes and c["rank"] == 1]
        
        if len(tied_candidates) > 1:
            ties.append({
                "election_id": item["election_id"],
                "election_title": item["election_title"],
                "position_id": item["position_id"],
                "position_name": item["position_name"],
                "tied_candidates": tied_candidates,
                "vote_count": max_votes,
                "status": "tie",
                "total_candidates": len(candidates),
            })
    
    return ties


def generate_dashboard_statistics(election_id=None, position_id=None):
    results_data = get_results_by_position(election_id=election_id, position_id=position_id)
    participation = calculate_participation(election_id=election_id)
    winners = determine_winners(election_id=election_id, position_id=position_id)
    ties = detect_ties(election_id=election_id, position_id=position_id)
    total_votes = sum(item["total_votes"] for item in results_data)

    if election_id:
        total_candidates = Candidate.query.filter_by(election_id=election_id).count()
        total_positions = Position.query.filter_by(election_id=election_id).count()
        total_elections = 1 if Election.query.filter_by(id=election_id).first() else 0
    else:
        total_candidates = Candidate.query.count()
        total_positions = Position.query.count()
        total_elections = Election.query.count()

    return {
        "total_votes": total_votes,
        "total_candidates": total_candidates,
        "total_positions": total_positions,
        "total_elections": total_elections,
        "participation": participation[0] if participation else {
            "registered_students": 0,
            "students_voted": 0,
            "students_remaining": 0,
            "participation_rate": 0.0,
        },
        "winners": winners,
        "results": results_data,
        "ties": ties,
    }


def get_live_statistics(election_id=None):
    election_filter = Election.query
    
    active_elections = election_filter.filter(Election.status == "Active").all()
    
    live_stats = []
    for election in active_elections:
        positions = Position.query.filter_by(election_id=election.id).all()
        for position in positions:
            candidates = Candidate.query.filter_by(position_id=position.id).all()
            candidate_stats = []
            total_votes = 0
            
            for candidate in candidates:
                vote_count = Vote.query.filter(
                    Vote.candidate_id == candidate.id,
                    Vote.election_id == election.id
                ).count()
                total_votes += vote_count
                candidate_stats.append({
                    "candidate_id": candidate.id,
                    "candidate_name": candidate.full_name or candidate.name,
                    "votes": vote_count,
                    "percentage": calculate_percentages(vote_count, total_votes) if total_votes else 0.0,
                })
            
            candidate_stats.sort(key=lambda x: -x["votes"])
            for rank, stat in enumerate(candidate_stats, start=1):
                stat["rank"] = rank
            
            live_stats.append({
                "election_id": election.id,
                "election_title": election.title,
                "position_id": position.id,
                "position_name": position.name,
                "candidates": candidate_stats,
                "total_votes": total_votes,
                "last_updated": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"),
                "status": "Active",
            })
    
    return live_stats


def get_candidate_statistics(election_id=None, candidate_id=None):
    candidate_filter = Candidate.query
    
    if election_id:
        candidate_filter = candidate_filter.filter(Candidate.election_id == election_id)
    if candidate_id:
        candidate_filter = candidate_filter.filter(Candidate.id == candidate_id)
    
    candidates = candidate_filter.all()
    stats = []
    
    for candidate in candidates:
        position = candidate.position
        election = candidate.election
        
        vote_count = Vote.query.filter(
            Vote.candidate_id == candidate.id,
            Vote.election_id == (election_id or candidate.election_id)
        ).count()
        
        position_candidates = Candidate.query.filter_by(position_id=position.id).all()
        total_position_votes = sum(
            Vote.query.filter(Vote.candidate_id == c.id).count()
            for c in position_candidates
        )
        
        stats.append({
            "candidate_id": candidate.id,
            "candidate_name": candidate.full_name or candidate.name,
            "candidate_photo": candidate.photo,
            "position_id": position.id,
            "position_name": position.name,
            "election_id": election.id,
            "election_title": election.title,
            "votes": vote_count,
            "percentage": calculate_percentages(vote_count, total_position_votes),
            "rank": None,
            "manifesto": candidate.manifesto,
            "biography": candidate.biography,
            "department": candidate.student.department if hasattr(candidate, 'student') and candidate.student else None,
            "class": candidate.student.class_name if hasattr(candidate, 'student') and candidate.student else None,
        })
    
    if election_id:
        results = get_results_by_position(election_id=election_id)
        for result in results:
            for candidate_stat in stats:
                for ranked_candidate in result.get("candidates", []):
                    if ranked_candidate["candidate_id"] == candidate_stat["candidate_id"]:
                        candidate_stat["rank"] = ranked_candidate["rank"]
                        break
    
    return stats


def get_position_statistics(election_id=None, position_id=None):
    position_filter = Position.query
    
    if election_id:
        position_filter = position_filter.filter(Position.election_id == election_id)
    if position_id:
        position_filter = position_filter.filter(Position.id == position_id)
    
    positions = position_filter.order_by(Position.display_order.asc(), Position.id.asc()).all()
    stats = []
    
    for position in positions:
        candidates = Candidate.query.filter_by(position_id=position.id).all()
        vote_counts = []
        total_votes = 0
        
        for candidate in candidates:
            votes = Vote.query.filter(
                Vote.candidate_id == candidate.id
            ).count()
            total_votes += votes
            vote_counts.append({
                "candidate_id": candidate.id,
                "votes": votes,
                "percentage": calculate_percentages(votes, total_votes),
            })
        
        vote_counts.sort(key=lambda x: -x["votes"])
        for rank, vc in enumerate(vote_counts, start=1):
            vc["rank"] = rank
        
        winner_candidate = candidates[0] if candidates else None
        winner_votes = vote_counts[0]["votes"] if vote_counts else 0
        
        stats.append({
            "position_id": position.id,
            "position_name": position.name,
            "position_code": position.position_code if hasattr(position, 'position_code') else None,
            "election_id": position.election_id,
            "election_title": position.election.title if position.election else None,
            "total_votes": total_votes,
            "winner": {
                "candidate_name": winner_candidate.full_name if winner_candidate else None,
                "votes": winner_votes,
                "percentage": vote_counts[0]["percentage"] if vote_counts else 0,
            },
            "candidates_count": len(candidates),
            "vote_distribution": vote_counts,
        })
    
    return stats


def get_election_summary(election_id):
    election = Election.query.get(election_id)
    if not election:
        return None
    
    results_data = get_results_by_position(election_id=election_id)
    participation = calculate_participation(election_id=election_id)
    winners = determine_winners(election_id=election_id)
    
    return {
        "election_name": election.title,
        "election_year": election.election_year,
        "election_status": election.status,
        "start_date": election.start_date,
        "end_date": election.end_date,
        "winner_count": len([w for w in winners if w["winner"]]),
        "positions_completed": len(results_data),
        "total_votes": sum(r["total_votes"] for r in results_data),
        "participation": participation[0] if participation else None,
        "winners": winners,
        "results": results_data,
    }


def verify_result_integrity(election_id):
    election = Election.query.get(election_id)
    if not election:
        return {"valid": False, "errors": ["Election not found"]}
    
    errors = []
    warnings = []
    
    positions = Position.query.filter_by(election_id=election_id).all()
    if not positions:
        errors.append("No positions configured for this election")
    
    candidates = Candidate.query.filter_by(election_id=election_id).all()
    if not candidates:
        errors.append("No candidates registered for this election")
    
    for position in positions:
        pos_candidates = Candidate.query.filter_by(position_id=position.id).all()
        if not pos_candidates:
            warnings.append(f"Position '{position.name}' has no candidates")
    
    total_votes = Vote.query.filter_by(election_id=election_id).count()
    if total_votes == 0:
        warnings.append("No votes recorded for this election")
    
    for candidate in candidates:
        votes = Vote.query.filter_by(candidate_id=candidate.id, election_id=election_id).count()
        if votes < 0:
            errors.append(f"Negative vote count for candidate {candidate.name}")
    
    unique_voters = db.session.query(Vote.student_id).filter(
        Vote.election_id == election_id
    ).distinct().count()
    
    total_student_votes = Vote.query.filter_by(
        election_id=election_id
    ).count()
    
    if unique_voters != total_student_votes:
        avg_votes_per_student = total_student_votes / unique_voters if unique_voters > 0 else 0
        if avg_votes_per_student > 1:
            warnings.append(f"Some students may have voted multiple times (avg {avg_votes_per_student:.2f} votes/student)")
    
    return {
        "valid": len(errors) == 0,
        "election_id": election_id,
        "election_title": election.title,
        "errors": errors,
        "warnings": warnings,
        "total_votes": total_votes,
        "total_candidates": len(candidates),
        "total_positions": len(positions),
    }


def publish_results(election_id):
    election = Election.query.get_or_404(election_id)
    election.visibility = "Visible"
    election.show_live_results = True
    db.session.add(election)
    db.session.add(AuditLog(user=str(session.get("admin_id")), action=f"Published results for {election.title}"))
    db.session.commit()
    return True


def unpublish_results(election_id):
    election = Election.query.get_or_404(election_id)
    election.visibility = "Hidden"
    election.show_live_results = False
    db.session.add(election)
    db.session.add(AuditLog(user=str(session.get("admin_id")), action=f"Hidden results for {election.title}"))
    db.session.commit()
    return True


def archive_results(election_id):
    election = Election.query.get_or_404(election_id)
    results_data = get_results_by_position(election_id=election_id)
    
    archive_entry = Archive(
        election_name=election.title,
        results=json.dumps(results_data),
        created_at=datetime.utcnow()
    )
    db.session.add(archive_entry)
    db.session.add(AuditLog(user=str(session.get("admin_id")), action=f"Archived results for {election.title}"))
    db.session.commit()
    
    return {"election_id": election.id, "election_title": election.title, "results": results_data}


def export_pdf(election_id):
    try:
        from fpdf import FPDF
    except ImportError:
        flash("PDF export requires fpdf library", "danger")
        return None
    
    election = Election.query.get(election_id)
    if not election:
        flash("Election not found", "danger")
        return None
    
    results_data = get_results_by_position(election_id=election_id)
    participation = calculate_participation(election_id=election_id)
    
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, f"Election Results: {election.title}", ln=True, align="C")
    
    if election.election_year:
        pdf.set_font("Arial", "", 12)
        pdf.cell(0, 10, f"Year: {election.election_year}", ln=True, align="C")
    
    pdf.ln(10)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Summary Statistics", ln=True)
    pdf.set_font("Arial", "", 12)
    
    part = participation[0] if participation else {}
    pdf.cell(0, 8, f"Registered Students: {part.get('registered_students', 0)}", ln=True)
    pdf.cell(0, 8, f"Students Voted: {part.get('students_voted', 0)}", ln=True)
    pdf.cell(0, 8, f"Participation Rate: {part.get('participation_rate', 0)}%", ln=True)
    
    for result in results_data:
        pdf.ln(5)
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, f"Position: {result['position_name']}", ln=True)
        pdf.set_font("Arial", "", 12)
        
        for candidate in result.get("candidates", []):
            winner_tag = " [WINNER]" if candidate["rank"] == 1 else ""
            pdf.cell(0, 8, f"  {candidate['candidate_name']}: {candidate['votes']} votes ({candidate['percentage']}%){winner_tag}", ln=True)
    
    db.session.add(AuditLog(user="ADMIN", action=f"Exported PDF for {election.title}"))
    db.session.commit()
    
    return pdf

def export_excel(election_id):
    try:
        import openpyxl
    except ImportError:
        flash("Excel export requires openpyxl library", "danger")
        return None
    
    election = Election.query.get(election_id)
    if not election:
        flash("Election not found", "danger")
        return None
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Election Results"
    
    ws["A1"] = f"Election: {election.title}"
    ws["A2"] = f"Year: {election.election_year}"
    ws["A3"] = f"Status: {election.status}"
    
    row = 5
    ws[f"A{row}"] = "Position"
    ws[f"B{row}"] = "Candidate"
    ws[f"C{row}"] = "Votes"
    ws[f"D{row}"] = "Percentage"
    ws[f"E{row}"] = "Rank"
    ws[f"F{row}"] = "Status"
    row += 1
    
    results_data = get_results_by_position(election_id=election_id)
    for result in results_data:
        for candidate in result.get("candidates", []):
            ws[f"A{row}"] = result["position_name"]
            ws[f"B{row}"] = candidate["candidate_name"]
            ws[f"C{row}"] = candidate["votes"]
            ws[f"D{row}"] = candidate["percentage"]
            ws[f"E{row}"] = candidate["rank"]
            ws[f"F{row}"] = "Winner" if candidate["rank"] == 1 else "Candidate"
            row += 1
    
    db.session.add(AuditLog(user="ADMIN", action=f"Exported Excel for {election.title}"))
    db.session.commit()
    
    return wb


def export_csv(election_id):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Election", "Position", "Candidate", "Votes", "Percentage", "Rank", "Status"])

    election = Election.query.get(election_id)
    if not election:
        flash("Election not found", "danger")
        return None
    
    results_data = get_results_by_position(election_id=election_id)
    for result in results_data:
        for candidate in result.get("candidates", []):
            writer.writerow([
                election.title,
                result["position_name"],
                candidate["candidate_name"],
                candidate["votes"],
                candidate["percentage"],
                candidate["rank"],
                "Winner" if candidate["rank"] == 1 else "Candidate"
            ])

    db.session.add(AuditLog(user="ADMIN", action=f"Exported CSV for {election.title}"))
    db.session.commit()

    response = Response(output.getvalue(), mimetype="text/csv")
    response.headers["Content-Disposition"] = f"attachment; filename={election.title.replace(' ', '_')}_results.csv"
    return response


def search_results(query, election_id=None, position=None, candidate=None, department=None, class_name=None, year=None, status=None):
    results_data = get_results_by_position(election_id=election_id)
    
    if query:
        query_lower = query.lower()
        filtered_results = []
        for r in results_data:
            if query_lower in r["position_name"].lower():
                filtered_results.append(r)
            else:
                filtered_candidates = [
                    c for c in r.get("candidates", [])
                    if query_lower in c["candidate_name"].lower()
                ]
                if filtered_candidates:
                    r_copy = r.copy()
                    r_copy["candidates"] = filtered_candidates
                    filtered_results.append(r_copy)
        results_data = filtered_results
    
    if position:
        results_data = [r for r in results_data if position.lower() in r["position_name"].lower()]
    
    if candidate:
        filtered_results = []
        for r in results_data:
            filtered_candidates = [
                c for c in r.get("candidates", [])
                if candidate.lower() in c["candidate_name"].lower()
            ]
            if filtered_candidates:
                r_copy = r.copy()
                r_copy["candidates"] = filtered_candidates
                filtered_results.append(r_copy)
        results_data = filtered_results
    
    if department:
        results_data = [
            r for r in results_data
            if any(
                (c.get("candidate_id") and 
                 Candidate.query.get(c["candidate_id"]).student and
                 Candidate.query.get(c["candidate_id"]).student.department and
                 department.lower() in Candidate.query.get(c["candidate_id"]).student.department.lower())
                for c in r.get("candidates", [])
            )
        ]
    
    if year:
        results_data = [r for r in results_data if r.get("election_year") == year]
    
    if status:
        results_data = [r for r in results_data if r.get("election_status", "").lower() == status.lower()]
    
    return results_data


def get_election_archive(election_id):
    archive = Archive.query.filter_by(election_name=Election.query.get(election_id).title).order_by(Archive.id.desc()).first()
    if archive:
        return json.loads(archive.results)
    return None


@results.route("/admin/results")
@admin_required
def results_page():
    election_id = request.args.get("election_id", type=int)
    position_filter = request.args.get("position", "")
    search_query = request.args.get("search", "")
    
    stats = generate_dashboard_statistics(election_id=election_id)
    elections = Election.query.order_by(Election.id.desc()).all()
    
    if search_query:
        stats["results"] = search_results(
            search_query,
            election_id=election_id,
            position=position_filter
        )
    
    ties = detect_ties(election_id=election_id)
    
    return render_template(
        "results_management.html",
        stats=stats,
        elections=elections,
        selected_election_id=election_id,
        results=stats["results"],
        ties=ties,
        search_query=search_query
    )


@results.route("/admin/results/live")
@admin_required
def live_results_page():
    election_id = request.args.get("election_id", type=int, default=None)
    
    stats = get_live_statistics(election_id=election_id)
    elections = Election.query.filter(Election.status == "Active").all()
    
    return render_template(
        "results_live.html",
        stats=stats,
        elections=elections,
        selected_election_id=election_id
    )


@results.route("/results/<int:election_id>")
def public_results(election_id):
    election = Election.query.get(election_id)
    
    if not election:
        flash("Election not found", "error")
        return redirect(url_for("index"))
    
    if election.status not in ["Closed", "Archived"]:
        flash("Results are not yet available for this election", "warning")
        return redirect(url_for("index"))
    
    if not election.show_live_results and election.visibility != "Visible":
        flash("Results are not publicly available", "error")
        return redirect(url_for("index"))
    
    results_data = get_results_by_position(election_id=election_id)
    participation = calculate_participation(election_id=election_id)
    winners = determine_winners(election_id=election_id)
    election_summary = get_election_summary(election_id)
    
    return render_template(
        "voter_results.html",
        election=election,
        results=results_data,
        participation=participation[0] if participation else None,
        winners=winners,
        election_summary=election_summary
    )


@results.route("/admin/results/winner/<int:position_id>")
@admin_required
def get_position_winner(position_id):
    position = Position.query.get_or_404(position_id)
    results_data = get_results_by_position(position_id=position_id)
    
    if not results_data:
        return jsonify({"winner": None, "votes": 0})
    
    result = results_data[0]
    winner = result.get("winner")
    
    return jsonify({
        "winner": winner.get("candidate_name") if winner else None,
        "votes": winner.get("votes") if winner else 0,
        "percentage": winner.get("percentage") if winner else 0,
    })


@results.route("/admin/results/integrity/<int:election_id>")
@admin_required
def check_integrity(election_id):
    integrity_report = verify_result_integrity(election_id)
    return jsonify(integrity_report)


@results.route("/admin/results/export/csv/<int:election_id>")
@admin_required
def export_csv_route(election_id):
    return export_csv(election_id)


@results.route("/admin/results/export/excel/<int:election_id>")
@admin_required
def export_excel_route(election_id):
    wb = export_excel(election_id)
    if wb:
        from openpyxl import Workbook
        import io
        from flask import Response
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = Response(
            output.read(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response.headers["Content-Disposition"] = f"attachment; filename={Election.query.get(election_id).title.replace(' ', '_')}_results.xlsx"
        return response
    return redirect(url_for("results.results_page"))


@results.route("/admin/results/export/pdf/<int:election_id>")
@admin_required
def export_pdf_route(election_id):
    pdf = export_pdf(election_id)
    if pdf:
        import io
        from flask import Response
        
        output = io.BytesIO()
        pdf.output(output, "F")
        output.seek(0)
        
        response = Response(
            output.read(),
            mimetype="application/pdf"
        )
        response.headers["Content-Disposition"] = f"attachment; filename={Election.query.get(election_id).title.replace(' ', '_')}_results.pdf"
        return response
    return redirect(url_for("results.results_page"))


@results.route("/admin/results/publish/<int:election_id>", methods=["POST"])
@admin_required
def publish_results_route(election_id):
    try:
        publish_results(election_id)
        flash("Results published successfully", "success")
    except Exception as e:
        flash(f"Error publishing results: {str(e)}", "danger")
    return redirect(url_for("results.results_page", election_id=election_id))


@results.route("/admin/results/unpublish/<int:election_id>", methods=["POST"])
@admin_required
def unpublish_results_route(election_id):
    try:
        unpublish_results(election_id)
        flash("Results hidden successfully", "success")
    except Exception as e:
        flash(f"Error unpublishing results: {str(e)}", "danger")
    return redirect(url_for("results.results_page", election_id=election_id))


@results.route("/admin/results/archive/<int:election_id>", methods=["POST"])
@admin_required
def archive_results_route(election_id):
    try:
        archive_results(election_id)
        flash("Election results archived successfully", "success")
    except Exception as e:
        flash(f"Error archiving results: {str(e)}", "danger")
    return redirect(url_for("results.results_page", election_id=election_id))


@results.route("/admin/results/print/<int:election_id>")
@admin_required
def print_results(election_id):
    election = Election.query.get(election_id)
    if not election:
        flash("Election not found", "danger")
        return redirect(url_for("results.results_page"))
    
    results_data = get_results_by_position(election_id=election_id)
    participation = calculate_participation(election_id=election_id)
    
    return render_template(
        "results_print.html",
        election=election,
        results=results_data,
        participation=participation[0] if participation else None
    )